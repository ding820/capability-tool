# app.py
import streamlit as st
import pandas as pd
import os, tempfile, shutil
import openpyxl
from data_loader import load_roster, load_expert_orders, load_manager_data, find_manager_cols
from rules import identify_experts, identify_managers
from history import save_period, get_previous_ids
from exporter import build_output_excel

HISTORY_PATH = os.path.join(os.path.dirname(__file__), 'history.json')

st.set_page_config(page_title='能力提升计划识别工具', layout='wide')
st.title('能力提升计划识别工具')

# ── 侧边栏：参数输入 ──────────────────────────────────────
with st.sidebar:
    st.header('识别参数')

    st.subheader('专家识别（双月）')
    month1 = st.text_input('第一个月（如 5月）', value='5月')
    month2 = st.text_input('第二个月（如 6月）', value='6月')
    expert_period_end = st.date_input('识别周期结束日', value=pd.Timestamp('2026-06-30'))

    st.subheader('主管识别（季度）')
    quarter = st.selectbox('季度', ['Q1', 'Q2', 'Q3', 'Q4'], index=1)
    year_str = st.text_input('年份后两位（如 26）', value='26')

    st.subheader('输出参数')
    period_label = st.text_input('输出表期间标签（如 7月）', value='7月')
    improve_period = st.text_input('改进周期（如 7月-8月）', value='7月-8月')

# ── 主体：文件上传 ────────────────────────────────────────
col1, col2, col3 = st.columns(3)
with col1:
    roster_file = st.file_uploader('花名册（需含"花名册最新" Sheet）', type=['xlsx'])
with col2:
    expert_file = st.file_uploader('专家定单底表（需含"专家底表" Sheet）', type=['xlsx'])
with col3:
    manager_file = st.file_uploader('主管底表（需含"主管底表" Sheet）', type=['xlsx'])

# ── 生成按钮 ──────────────────────────────────────────────
if st.button('生成能力提升名单', type='primary'):
    if not (roster_file and expert_file and manager_file):
        st.error('请上传全部3个文件')
        st.stop()

    with st.spinner('计算中...'):
        not_graduated, new_entries, triggered_ids = [], [], set()
        tmp_dir = tempfile.mkdtemp()
        try:
            roster_path = os.path.join(tmp_dir, 'roster.xlsx')
            expert_path = os.path.join(tmp_dir, 'expert.xlsx')
            manager_path = os.path.join(tmp_dir, 'manager.xlsx')
            for src, dst in [(roster_file, roster_path), (expert_file, expert_path), (manager_file, manager_path)]:
                with open(dst, 'wb') as f:
                    f.write(src.read())

            try:
                roster = load_roster(roster_path)
            except Exception as e:
                st.error(f'花名册读取失败：{e}\n请确认文件包含"花名册最新" Sheet，且"雇佣状态"列存在。')
                st.stop()

            try:
                expert_orders = load_expert_orders(expert_path)
            except Exception as e:
                st.error(f'专家底表读取失败：{e}\n请确认文件包含"专家底表" Sheet。')
                st.stop()

            try:
                manager_raw = load_manager_data(manager_path)
            except Exception as e:
                st.error(f'主管底表读取失败：{e}\n请确认文件包含"主管底表" Sheet。')
                st.stop()

            try:
                # 合并专家数据（花名册补充门店/职级/入职日期），只取岗位类别=='产品专家'的人
                expert_merged = expert_orders.merge(
                    roster[roster['岗位类别'] == '产品专家'][
                        ['工号', '员工姓名', '二级部门', '三级部门', '四级部门', '新岗位名称', '新职级', '入职日期', '岗位类别']
                    ],
                    left_on='专家工号', right_on='工号', how='inner'
                ).drop(columns=['专家工号'])

                # 合并主管数据，并将 '姓名' rename 为 '员工姓名'
                manager_merged = manager_raw.rename(columns={'姓名': '员工姓名'}).merge(
                    roster[['工号', '二级部门', '三级部门', '四级部门', '新岗位名称', '新职级', '入职日期']],
                    on='工号', how='inner'
                )

                # 识别专家
                period_end_ts = pd.Timestamp(expert_period_end)
                expert_result = identify_experts(expert_merged, period_end_ts)

                # 识别主管
                achieve_cols, target_cols = find_manager_cols(list(manager_merged.columns), quarter, year_str)
                if not achieve_cols or not target_cols:
                    st.error(
                        f'未在主管底表中找到 {quarter} 对应的达成/目标列。\n'
                        f'请检查列名是否包含"{year_str}年X月"或"20{year_str}年X月"格式。'
                    )
                    st.stop()
                manager_result = identify_managers(manager_merged, achieve_cols, target_cols, period_end_ts)

                # 汇总触发人员
                expert_triggered = expert_result[expert_result['触发识别']].to_dict('records')
                manager_triggered = manager_result[manager_result['触发识别']].to_dict('records')
                all_triggered = expert_triggered + manager_triggered
                triggered_ids = {r['工号'] for r in all_triggered}

                # 区分未出营 vs 新进入
                prev_ids = get_previous_ids(HISTORY_PATH)
                not_graduated = [r for r in all_triggered if r['工号'] in prev_ids]
                new_entries = [r for r in all_triggered if r['工号'] not in prev_ids]
            except Exception as e:
                st.error(f'计算识别结果时出错：{e}')
                st.stop()

        finally:
            shutil.rmtree(tmp_dir)

    st.success(f'识别完成：未出营 {len(not_graduated)} 人，新进入 {len(new_entries)} 人')

    col_a, col_b = st.columns(2)
    with col_a:
        st.subheader(f'未出营名单（{len(not_graduated)} 人）')
        if not_graduated:
            st.dataframe(pd.DataFrame(not_graduated)[['工号', '员工姓名', '四级部门', '新岗位名称', '新职级', '识别原因']])
        else:
            st.info('无未出营人员')

    with col_b:
        st.subheader(f'新进入名单（{len(new_entries)} 人）')
        if new_entries:
            st.dataframe(pd.DataFrame(new_entries)[['工号', '员工姓名', '四级部门', '新岗位名称', '新职级', '识别原因']])
        else:
            st.info('无新进入人员')

    excel_buf = build_output_excel(period_label, improve_period, not_graduated, new_entries)
    st.download_button(
        label='下载公示表 Excel',
        data=excel_buf,
        file_name=f'{period_label}能力提升计划最终名单-公示.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )

    period_key = f'{month1}/{month2}'
    save_period(HISTORY_PATH, period_key, list(triggered_ids))
    st.caption(f'本期名单已保存至历史记录（周期：{period_key}）')

# ── 历史初始化（可选） ────────────────────────────────────
with st.expander('初始化历史记录（首次使用时导入上期名单）'):
    st.write('上传已有公示表 Excel，从中提取上期进入名单工号作为历史基准。')
    hist_file = st.file_uploader('上传历史公示表', type=['xlsx'], key='hist')
    hist_period = st.text_input('历史周期标签（如 5月/6月）', key='hist_period')
    if st.button('导入历史'):
        if hist_file and hist_period:
            with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
                f.write(hist_file.read())
                tmp = f.name
            try:
                wb = openpyxl.load_workbook(tmp)
            finally:
                os.unlink(tmp)
            ids = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row_data in ws.iter_rows(values_only=True):
                    val = row_data[4] if len(row_data) > 4 else None
                    if val is not None:
                        s = str(val).strip()
                        # 工号至少4位纯数字
                        if s.isdigit() and len(s) >= 4:
                            ids.append(str(int(s)))
            save_period(HISTORY_PATH, hist_period, list(set(ids)))
            st.success(f'已导入 {len(set(ids))} 条历史记录（周期：{hist_period}）')
        else:
            st.warning('请上传文件并填写周期标签')
