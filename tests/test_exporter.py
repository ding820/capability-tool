# tests/test_exporter.py
import sys
import os
import io

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
from exporter import build_output_excel


def _sample_records():
    return [
        {
            '工号': '1001', '员工姓名': '张三', '二级部门': '湖北战区',
            '三级部门': '湖北一部', '四级部门': '武汉龙阳大道零售中心',
            '新岗位名称': '产品专家', '新职级': 13,
            '识别原因': '双月专家定单量排名在所在门店的后15%',
        },
        {
            '工号': '1002', '员工姓名': '李四', '二级部门': '湖北战区',
            '三级部门': '湖北三部', '四级部门': '恩施金桂大道汽车城零售中心',
            '新岗位名称': '高级产品专家', '新职级': 14,
            '识别原因': '双月高级专家定单量排名在省区专家岗位后50%',
        },
    ]


def test_build_output_excel_returns_bytesio():
    buf = build_output_excel('7月', '7月-8月', [_sample_records()[0]], [_sample_records()[1]])
    assert isinstance(buf, io.BytesIO)
    assert buf.tell() == 0
    assert len(buf.getvalue()) > 0


def test_output_excel_has_correct_sheet():
    buf = build_output_excel('7月', '7月-8月', [], _sample_records())
    wb = openpyxl.load_workbook(buf)
    assert '7月能力提升计划最终名单-公示' in wb.sheetnames


def test_output_contains_names():
    buf = build_output_excel('7月', '7月-8月', [], _sample_records())
    wb = openpyxl.load_workbook(buf)
    ws = wb['7月能力提升计划最终名单-公示']
    flat = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
    assert '张三' in flat
    assert '李四' in flat


def test_not_graduated_shows_no():
    buf = build_output_excel('7月', '7月-8月', _sample_records(), [])
    wb = openpyxl.load_workbook(buf)
    ws = wb['7月能力提升计划最终名单-公示']
    flat = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
    assert '否' in flat


def test_new_entries_shows_yes_and_period():
    buf = build_output_excel('7月', '7月-8月', [], _sample_records())
    wb = openpyxl.load_workbook(buf)
    ws = wb['7月能力提升计划最终名单-公示']
    flat = [cell.value for row in ws.iter_rows() for cell in row if cell.value]
    assert '是' in flat
    assert '7月-8月' in flat


def test_both_lists_empty_does_not_raise():
    buf = build_output_excel('7月', '7月-8月', [], [])
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    assert ws is not None


def test_sequence_numbers_start_at_1():
    buf = build_output_excel('7月', '7月-8月', _sample_records(), [])
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    # Row 1 = section title, row 2 = headers, row 3 = first data row
    assert ws.cell(row=3, column=1).value == 1
    assert ws.cell(row=4, column=1).value == 2
